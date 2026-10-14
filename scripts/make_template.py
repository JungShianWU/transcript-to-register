#!/usr/bin/env python3
"""從零產生通用範本 template/generic.xlsx（地號清冊格式）。

用法：python3 scripts/make_template.py   （需要 openpyxl）

欄位位置與常見的地號清冊一致（資料從第 5 列開始），
因此 src/core.js 可以同時填寫本範本或公司自有範本：
  第 5 列：區塊首列樣式＋各欄公式
  第 6 列：區塊中間列樣式
  第 11 列：區塊末列樣式
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, column_index_from_string as ci
from openpyxl.workbook.defined_name import DefinedName

OUT = Path(__file__).resolve().parent.parent / 'template' / 'generic.xlsx'
AUTHOR = 'Jacky'
FONT = 'Microsoft JhengHei'

# (欄, 第3列標題, 第4列單位, 欄寬, 數字格式)
COLS = [
    ('A', '區域', '', 8, None), ('B', '地段', '', 8, None), ('C', '小段', '', 8, None),
    ('D', '地號', '', 8, None), ('E', '面積', '(㎡)', 9, '#,##0.00'),
    ('F', '所有權人', '', 10, None), ('G', '統一編號', '', 12, None),
    ('H', '持分', '分子', 7, '0'), ('I', '', '/', 2, None), ('J', '', '分母', 7, '0'),
    ('K', '容積率', '(%)', 7, '0%'), ('L', '持有面積', '(坪)', 9, '#,##0.00'),
    ('M', '獎勵值', '(%)', 7, '0%'), ('N', '允建容積', '(坪)', 9, '#,##0.00'),
    ('O', '公告現值', '(元/㎡)', 10, '#,##0'),
    ('P', '建號', '', 7, None), ('Q', '總面積', '(㎡)', 8, '#,##0.00'),
    ('R', '主建物', '(㎡)', 8, '#,##0.00'),
    ('S', '持分', '', 4, '0'), ('T', '', '/', 2, None), ('U', '', '', 4, '0'),
    ('V', '主建物', '(坪)', 8, '#,##0.00'),
    ('W', '騎樓', '(㎡)', 7, '#,##0.00'),
    ('X', '持分', '', 4, '0'), ('Y', '', '/', 2, None), ('Z', '', '', 4, '0'),
    ('AA', '騎樓', '(坪)', 7, '#,##0.00'),
    ('AB', '附屬建物', '(㎡)', 8, '#,##0.00'),
    ('AC', '持分', '', 4, '0'), ('AD', '', '/', 2, None), ('AE', '', '', 4, '0'),
    ('AF', '附屬建物', '(坪)', 8, '#,##0.00'),
    ('AG', '平台', '(㎡)', 7, '#,##0.00'), ('AH', '平台', '(坪)', 7, '#,##0.00'),
    ('AI', '夾層', '(㎡)', 7, '#,##0.00'), ('AJ', '夾層', '(坪)', 7, '#,##0.00'),
    ('AK', '屋突', '(㎡)', 7, '#,##0.00'), ('AL', '屋突', '(坪)', 7, '#,##0.00'),
    ('AM', '防空避難室', '(㎡)', 8, '#,##0.00'), ('AN', '防空避難室', '(坪)', 8, '#,##0.00'),
    ('AO', '共有部分\n(一)', '(㎡)', 8, '#,##0.00'),
    ('AP', '持分', '分子', 6, '0'), ('AQ', '', '/', 2, None), ('AR', '', '分母', 6, '0'),
    ('AS', '持有面積', '(㎡)', 8, '#,##0.00'), ('AT', '持有面積', '(坪)', 8, '#,##0.00'),
    ('AU', '共有部分\n(二)', '(㎡)', 8, '#,##0.00'),
    ('AV', '持分', '分子', 6, '0'), ('AW', '', '/', 2, None), ('AX', '', '分母', 6, '0'),
    ('AY', '持有面積', '(㎡)', 8, '#,##0.00'), ('AZ', '持有面積', '(坪)', 8, '#,##0.00'),
    ('BA', '建物總面積', '(坪)', 9, '#,##0.00'),
    ('BB', '樓次', '', 6, None), ('BC', '門牌', '', 30, None),
    ('BD', '立體價差', '', 7, '0.00'), ('BE', '權利價值', '', 9, '#,##0.00'),
    ('BF', '分配比例', '', 8, '0.00%'),
]
LAST = 'BF'

# 第 5 列公式（core.js 會把第 5 列的相對參照換成各列列號，$E$5 換成各區塊首列）
FORMULAS = {
    'K': '=$R$1', 'M': '=$U$1',
    'L': '=$E$5*0.3025*H5/J5', 'N': '=L5*K5*(1+M5)',
    'V': '=R5*0.3025', 'AA': '=W5*0.3025', 'AF': '=AB5*0.3025', 'AH': '=AG5*0.3025',
    'AJ': '=AI5*0.3025', 'AL': '=AK5*0.3025', 'AN': '=AM5*0.3025',
    'AS': '=AO5*AP5/AR5', 'AT': '=AS5*0.3025',
    'AY': '=AU5*AV5/AX5', 'AZ': '=AY5*0.3025',
    'BA': '=V5+AA5+AF5+AH5+AJ5+AL5+AN5+AT5+AZ5',
    'BE': '=N5*BD5',
    'BF': '=IF(SUM($BE$5:$BE$5000)=0,0,BE5/SUM($BE$5:$BE$5000))',
}

thin = Side(style='thin', color='9AA5A0')
med = Side(style='medium', color='2F5D50')
HEAD_FILL = {'land': 'E3EEE9', 'bldg': 'EAF0F7', 'value': 'F6EFE3'}
INPUT_FILL = PatternFill('solid', fgColor='FFF4C2')


def group_of(col):
    n = ci(col)
    if n <= ci('O'):
        return 'land'
    if n <= ci('BC'):
        return 'bldg'
    return 'value'


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = '地號清冊'
    base = Font(name=FONT, size=10)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # 第 1 列：案名 + 參數輸入
    ws['A1'] = '地號清冊'
    ws['A1'].font = Font(name=FONT, size=14, bold=True)
    ws.merge_cells('A1:O1')
    for ref, text in (('P1', '基準容積率'), ('S1', '容積獎勵')):
        ws[ref] = text
        ws[ref].font = Font(name=FONT, size=10, bold=True)
        ws[ref].alignment = Alignment(horizontal='right', vertical='center')
    ws.merge_cells('P1:Q1')
    ws.merge_cells('S1:T1')
    ws.merge_cells('U1:V1')
    for ref, fmt in (('R1', '0%'), ('U1', '0%')):
        c = ws[ref]
        c.fill = INPUT_FILL
        c.number_format = fmt
        c.font = Font(name=FONT, size=11, bold=True)
        c.alignment = center
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws['X1'] = '← 黃色格請依案件填入，例如容積率 300% 填 300%、獎勵 80% 填 80%'
    ws['X1'].font = Font(name=FONT, size=9, color='5D6762')

    # 第 2 列：區段標題
    for rng, text, g in (('A2:O2', '土地標示部', 'land'), ('P2:BC2', '建物標示部', 'bldg'), ('BD2:BF2', '權利價值', 'value')):
        ws.merge_cells(rng)
        c = ws[rng.split(':')[0]]
        c.value = text
        c.font = Font(name=FONT, size=11, bold=True)
        c.alignment = center
        c.fill = PatternFill('solid', fgColor=HEAD_FILL[g])

    # 第 3、4 列：欄名與單位
    for col, h3, h4, width, _ in COLS:
        ws.column_dimensions[col].width = width
        fill = PatternFill('solid', fgColor=HEAD_FILL[group_of(col)])
        for r, v in ((3, h3), (4, h4)):
            c = ws[f'{col}{r}']
            if v:
                c.value = v
            c.font = Font(name=FONT, size=10, bold=(r == 3))
            c.alignment = center
            c.fill = fill
            c.border = Border(left=thin, right=thin, top=thin, bottom=(med if r == 4 else thin))
    for rng in ('H3:J3', 'S3:U3', 'X3:Z3', 'AC3:AE3', 'AP3:AR3', 'AV3:AX3'):
        ws.merge_cells(rng)
    for col in ('A', 'B', 'C', 'D', 'F', 'G', 'P', 'BB', 'BC', 'BD', 'BE', 'BF'):
        ws.merge_cells(f'{col}3:{col}4')
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 40
    ws.row_dimensions[4].height = 18

    # 第 5～11 列：樣板區塊（首列 5、中間列 6～10、末列 11）
    fmt = {c[0]: c[4] for c in COLS}
    for r in range(5, 12):
        top = med if r == 5 else thin
        bottom = med if r == 11 else thin
        ws.row_dimensions[r].height = 22
        for i in range(1, ci(LAST) + 1):
            col = get_column_letter(i)
            c = ws[f'{col}{r}']
            c.font = base
            c.alignment = Alignment(horizontal='left' if col in ('F', 'BC') else 'center', vertical='center')
            c.border = Border(left=thin, right=thin, top=top, bottom=bottom)
            if fmt.get(col):
                c.number_format = fmt[col]
    for col, f in FORMULAS.items():
        ws[f'{col}5'] = f
    for col in ('I', 'T', 'AD', 'AQ', 'AW'):
        ws[f'{col}5'] = '/'

    ws.freeze_panes = 'H5'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:4'
    ws.oddFooter.center.text = '&P / &N'
    wb.defined_names.add(DefinedName('_xlnm.Print_Area', localSheetId=0, attr_text=f"'地號清冊'!$A$1:${LAST}$11"))

    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.title = '地號清冊'
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f'✓ {OUT}')


if __name__ == '__main__':
    main()
